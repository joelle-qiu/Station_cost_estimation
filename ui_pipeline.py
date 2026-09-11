# -*- coding: utf-8 -*-
"""筛站 / 配储评分两个页签。作者: jiali.qiu"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st

from intake_screening import (
    SCREEN_HIGHLIGHT_COLS,
    SCREEN_RESULT_FRONT,
    order_screen_columns,
    run_screen,
)
from tariff_lib import coverage_rows, load_library
from ui_scoring import render_scoring_page, reset_score_wizard

ROOT = Path(__file__).resolve().parent
OVERVIEW_HTML = ROOT / "doc" / "配储与VPP工具全览.html"
TEMPLATE_CANDIDATES = [
    ROOT / "doc" / "场站资源摸底收资表.xlsx",
    ROOT / "doc" / "场站资源摸底收资表（客户版）.xlsx",
    ROOT / "doc" / "通用场站配储摸底收资表.xlsx",
]


def _blank_template() -> Path | None:
    found = [p for p in TEMPLATE_CANDIDATES if p.exists() and not p.name.startswith("~$")]
    if not found:
        return None
    return max(found, key=lambda p: p.stat().st_mtime)


def _style_screen_df(df: pd.DataFrame):
    """结论、推荐下一步黄底加粗。"""
    cols = [c for c in SCREEN_HIGHLIGHT_COLS if c in df.columns]

    def _hl(_val: object) -> str:
        return "background-color: #FFF3B0; font-weight: 600; color: #1A1A1A"

    sty = df.style
    if cols:
        if hasattr(sty, "map"):
            sty = sty.map(_hl, subset=cols)
        else:
            sty = sty.applymap(_hl, subset=cols)
    if "原因" in df.columns:
        sty = sty.set_properties(
            subset=["原因"],
            **{"text-align": "left", "white-space": "pre-wrap"},
        )
    return sty


def _export_screen_xlsx(df: pd.DataFrame) -> bytes:
    """导出筛站结果：结论 / 推荐下一步黄底，原因自动换行。"""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ordered = order_screen_columns(df)
    buf = BytesIO()
    ordered.to_excel(buf, index=False, engine="openpyxl")
    buf.seek(0)
    wb = load_workbook(buf)
    ws = wb.active
    ws.title = "筛站结果"
    yellow = PatternFill("solid", fgColor="FFF3B0")
    wrap = Alignment(wrap_text=True, vertical="center", horizontal="left")
    header_font = Font(bold=True, color="1A1A1A")
    body_bold = Font(bold=True, color="1A1A1A")
    thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    widths = {
        "红绿灯": 10,
        "结论": 18,
        "推荐下一步": 44,
        "原因": 72,
        "场站名称": 22,
        "EGC预制": 30,
        "缺失": 22,
        "资源类型": 12,
        "分型倾向": 12,
    }
    headers = [cell.value for cell in ws[1]]
    for idx, name in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = widths.get(str(name), 14)
        highlight = name in SCREEN_HIGHLIGHT_COLS
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row, idx)
            cell.alignment = wrap
            cell.border = thin
            if row == 1:
                cell.font = header_font
            if highlight:
                cell.fill = yellow
                if row > 1:
                    cell.font = body_bold
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 80
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _download_template_btn() -> None:
    path = _blank_template()
    if path is not None:
        st.download_button(
            "下载空白摸底表（发给客户填写）",
            data=path.read_bytes(),
            file_name="场站资源摸底收资表.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    else:
        st.caption("尚未生成摸底表，请先运行 `_gen_intake_xlsx.py`。")


def _excel_buffer(up) -> BytesIO:
    if hasattr(up, "getvalue"):
        return BytesIO(up.getvalue())
    data = up.read()
    if hasattr(up, "seek"):
        try:
            up.seek(0)
        except (OSError, ValueError):
            pass
    return BytesIO(data)


def _read_station_sheet(up) -> pd.DataFrame:
    """读取上传的摸底表。只读一次字节，避免 Streamlit 文件指针被读空。"""
    name = str(getattr(up, "name", "") or "").lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        raise RuntimeError("请把文件另存为 .xlsx 后再上传。旧版 .xls 本工具不再读取。")
    buf = _excel_buffer(up)
    try:
        xls = pd.ExcelFile(buf, engine="openpyxl")
    except ImportError as exc:
        raise RuntimeError(
            "当前运行环境缺少 openpyxl，无法读取 Excel。"
            "请关闭本页面，在工具目录双击 install.bat，"
            "或在命令行执行：venv\\Scripts\\pip install openpyxl"
        ) from exc
    except Exception as exc:
        raise RuntimeError(f"无法打开 Excel，请确认文件未损坏、不是正在被 Excel 占用的临时文件。原因：{exc}") from exc

    sheet = "2-资源清单" if "2-资源清单" in xls.sheet_names else (
        "2-场站资产清单" if "2-场站资产清单" in xls.sheet_names else xls.sheet_names[-1]
    )
    peek = pd.read_excel(xls, sheet_name=sheet, header=None, nrows=8)
    header_row = 0
    for i in range(min(8, len(peek))):
        vals = [str(x).strip() for x in peek.iloc[i].tolist()]
        if "场站名称" in vals:
            header_row = i
            break
    raw = pd.read_excel(xls, sheet_name=sheet, header=header_row)
    raw = raw.dropna(how="all")
    name_col = "场站名称" if "场站名称" in raw.columns else raw.columns[0]
    raw = raw[raw[name_col].astype(str).str.strip().ne("")]
    raw = raw[~raw[name_col].astype(str).str.contains("示例|说明|请勿", na=False)]
    if raw.empty:
        raise RuntimeError("没有读到有效场站。请填写「2-资源清单」，一行一个场站，不要只留示例行。")
    return raw


def render_overview_page() -> None:
    """工具第一页：嵌入配储与 VPP 流程全览 HTML。"""
    st.title("流程说明 · 配储筛站与虚拟电厂组盘")
    st.caption(
        "这是整套工具怎么用、每一步输入什么、产出什么。看完点顶部「① 高价值筛站」开始。"
        "充电站和厂区/园区同一套流程；客户只填一张摸底表。"
    )
    go = st.button("开始① 高价值筛站", type="primary")
    if go:
        st.session_state.workbench_page = "① 高价值筛站"
        st.rerun()
    path = OVERVIEW_HTML
    if not path.exists():
        st.error("找不到 `doc/配储与VPP工具全览.html`。")
        return
    html = path.read_text(encoding="utf-8")
    import streamlit.components.v1 as components

    components.html(html, height=2200, scrolling=True)


def render_screening_page() -> None:
    st.subheader("① 高价值筛站（VPP 组盘 + 配储）")
    st.caption(
        "导入客户填好的统一摸底表（「2-资源清单」）。"
        "蓝列看 VPP 粗可调（含换电仓/柴发能否调度），绿列看储能投资红/黄/绿（变压器、日电量、场地；不问可调）。"
        "充电运营商和产业园/厂区同一套。不算配储 IRR。"
        "收益只按「电网代理购电 + 政府核定固定峰谷」测算；"
        "已入市标 ⭐，即将改规则标【政策即将变动需复核】。"
    )
    c1, c2 = st.columns([1, 2])
    with c1:
        _download_template_btn()
    up = st.file_uploader("上传填好的 Excel", type=["xlsx", "xls"])
    result = None
    if up is not None:
        try:
            raw = _read_station_sheet(up)
            result = run_screen(raw)
        except RuntimeError as exc:
            st.error(str(exc))
            return
        except Exception as exc:
            st.error(f"筛站计算失败：{exc}")
            return
        st.session_state["screen_result"] = result
    elif st.session_state.get("screen_result") is not None:
        result = st.session_state["screen_result"]
        st.caption("正在显示上次导入的筛站结果。重新上传可覆盖。切到②不会丢这份表。")
    else:
        st.info("还没有文件。先把 `doc/场站资源摸底收资表.xlsx` 发给客户（蓝列 VPP、绿列储能投资；充电站和厂区共用这一张）。")
        lib = load_library()
        with st.expander("当前电价库覆盖（测算用，须看标记）", expanded=True):
            st.markdown(
                f"**截至 {lib.get('as_of')}。** {lib.get('note', '')}\n\n"
                f"- {lib.get('legend', {}).get('star', '')}\n"
                f"- {lib.get('legend', {}).get('pending', '')}\n\n"
                "库内峰/谷**只给电网代理购电、政府核定固定峰谷**用。"
                "标了 ⭐ 的用户类型，页签②必须用手填零售/现货价，不能套表里的峰谷。"
            )
            st.dataframe(
                pd.DataFrame(coverage_rows(lib)),
                use_container_width=True,
                hide_index=True,
            )
        return

    st.session_state["screen_result"] = result

    vc = result["红绿灯"].value_counts()
    adj_sum = float(pd.to_numeric(result.get("粗可调kW"), errors="coerce").fillna(0).sum())
    m1, m2, m3, m4, m5 = st.columns(5)
    m1.metric("配储·绿", int(vc.get("绿", 0)))
    m2.metric("配储·黄", int(vc.get("黄", 0)))
    m3.metric("配储·红", int(vc.get("红", 0)))
    m4.metric("配储·灰存量", int(vc.get("灰", 0)))
    m5.metric("VPP 粗可调 kW", f"{adj_sum:.0f}")

    st.caption(
        "黄底两列是**结论**和**推荐下一步**，方便扫一眼。"
        "**原因**写清：为什么值或不值、对配储意味着什么、虚拟电厂是否还能组盘。"
    )
    view = result[[c for c in SCREEN_RESULT_FRONT if c in result.columns]]
    st.dataframe(
        _style_screen_df(view),
        use_container_width=True,
        hide_index=True,
        height=560,
        column_config={
            "结论": st.column_config.TextColumn("结论", width="small"),
            "推荐下一步": st.column_config.TextColumn("推荐下一步", width="medium"),
            "原因": st.column_config.TextColumn("原因", width="large"),
        },
    )

    try:
        xlsx_bytes = _export_screen_xlsx(result)
    except ImportError:
        st.warning("缺少 openpyxl，无法导出 Excel。请双击 install.bat 补装依赖。")
    else:
        st.download_button(
            "导出筛站结果（交给负责人）",
            data=xlsx_bytes,
            file_name="场站筛站结果_VPP与配储.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    names = result.loc[result["红绿灯"].isin(["绿", "黄"]), "场站名称"].astype(str).tolist()
    if names:
        pick = st.selectbox("把黄/绿站带入②评分（手填曲线结论）", ["—"] + names)
        go = st.button("带入并打开② 配储评分", type="primary", disabled=(pick == "—"))
        if pick != "—" and go:
            st.session_state["score_prefill"] = result[result["场站名称"].astype(str) == pick].iloc[0].to_dict()
            reset_score_wizard()
            st.session_state["workbench_page"] = "② 配储评分测算"
            st.rerun()
        elif pick != "—":
            st.session_state["score_prefill"] = result[result["场站名称"].astype(str) == pick].iloc[0].to_dict()
            st.caption("已记住该站电价与资产。点上面按钮，或点顶部「② 配储评分测算」。")
