# -*- coding: utf-8 -*-
"""筛站 / 评分 / 电价匹配冒烟。作者: jiali.qiu"""

from __future__ import annotations

import unittest
from pathlib import Path

import pandas as pd

from intake_screening import INTAKE_COLUMNS, STORAGE_FIRST_COL, STORAGE_LAST_COL, VPP_LAST_COL, run_screen, screen_station
from storage_scoring import ScoreInput, evaluate_hard_gate, score_and_invest
from tariff_lib import load_library, match_tariff, tariff_schedule_rows


def _row(**kwargs) -> pd.Series:
    rec = {c: "" for c in INTAKE_COLUMNS}
    rec.update(kwargs)
    return pd.Series(rec)


class TariffMatchTests(unittest.TestCase):
    def setUp(self) -> None:
        self.lib = load_library()

    def test_shenyang_is_liaoning(self) -> None:
        t = match_tariff("辽宁", "沈阳", self.lib)
        assert t is not None
        self.assertEqual(t["id"], "liaoning")
        self.assertTrue(t["has_tou"])

    def test_shenzhen_not_guangdong(self) -> None:
        t = match_tariff("广东", "深圳", self.lib)
        assert t is not None
        self.assertEqual(t["id"], "shenzhen")

    def test_guangzhou_is_guangdong(self) -> None:
        t = match_tariff("广东", "广州", self.lib)
        assert t is not None
        self.assertEqual(t["id"], "guangdong")
        self.assertTrue(t["market_no_gov_tou"])

    def test_liaoning_market_is_dynamic(self) -> None:
        t = match_tariff("辽宁", "沈阳", self.lib)
        assert t is not None
        self.assertTrue(t["market_no_gov_tou"])
        self.assertEqual(t["market"]["marker"], "⭐")
        self.assertTrue(t["agent"]["usable_for_spread"])
        self.assertFalse(t["market"]["usable_for_spread"])

    def test_schedule_table_has_clock_and_price(self) -> None:
        t = match_tariff("辽宁", "沈阳", self.lib)
        rows = tariff_schedule_rows(t, purchase_market=False)
        kinds = {r["类型"] for r in rows}
        self.assertIn("峰", kinds)
        self.assertIn("谷", kinds)
        peak = [r for r in rows if r["类型"] == "峰"]
        valley = [r for r in rows if r["类型"] == "谷"]
        self.assertGreaterEqual(len(peak), 2)
        self.assertGreaterEqual(len(valley), 2)
        self.assertTrue(any("07:30" in r["时段"] for r in peak))
        self.assertTrue(any("22:00" in r["时段"] or "05:00" in r["时段"] for r in valley))
        self.assertTrue(all("须手填" not in r["库内电价（元/kWh）"] for r in peak))

    def test_schedule_market_prices_must_be_filled(self) -> None:
        t = match_tariff("辽宁", "沈阳", self.lib)
        rows = tariff_schedule_rows(t, purchase_market=True)
        peak = [r for r in rows if r["类型"] == "峰"]
        self.assertTrue(any("须手填" in r["库内电价（元/kWh）"] for r in peak))

    def test_coverage_splits_agent_and_market(self) -> None:
        from tariff_lib import coverage_rows

        rows = coverage_rows(self.lib)
        marks = [r["标记"] for r in rows]
        self.assertTrue(any(m.startswith("⭐") for m in marks))
        self.assertTrue(any("即将变动" in (m or "") for m in marks))
        shandong = [r for r in rows if r["地区"] == "山东" and "代理购电" in r["适用用户"]]
        self.assertEqual(shandong[0]["标记"], "【政策即将变动需复核】")

    def test_intake_vpp_then_storage_split(self) -> None:
        self.assertEqual(INTAKE_COLUMNS[INTAKE_COLUMNS.index(VPP_LAST_COL) + 1], STORAGE_FIRST_COL)
        self.assertEqual(INTAKE_COLUMNS[INTAKE_COLUMNS.index(STORAGE_LAST_COL) + 1], "备注")
        self.assertEqual(STORAGE_FIRST_COL, "白天压不下去的负荷（千瓦）")
        self.assertLess(INTAKE_COLUMNS.index("换电仓功率能否远程下调"), INTAKE_COLUMNS.index(STORAGE_FIRST_COL))
        self.assertLess(INTAKE_COLUMNS.index("柴发"), INTAKE_COLUMNS.index(STORAGE_FIRST_COL))
        self.assertLess(INTAKE_COLUMNS.index("充电桩可远程限功率（千瓦）"), INTAKE_COLUMNS.index(STORAGE_FIRST_COL))


class ScreenTests(unittest.TestCase):
    def setUp(self) -> None:
        self.lib = load_library()
        self.ln = match_tariff("辽宁", "沈阳", self.lib)
        self.gd = match_tariff("广东", "广州", self.lib)

    def test_factory_without_piles_can_be_green(self) -> None:
        row = _row(
            场站名称="铁西机械厂",
            省="辽宁",
            市="沈阳",
            场景类型="厂区负荷",
            有厂区稳定负荷="是",
            厂区白天最小kW=380,
            变压器额定kVA=800,
            独立户号="是",
            共表或分表="分表独立户号",
            能调控="能",
            场地能否放储能柜="是",
            用电性质="工业",
            日均用电或充电量kWh=9000,
            购电方式="电网代理购电",
        )
        judged = screen_station(row, self.ln)
        self.assertIn(judged["红绿灯"], ("绿", "黄"))
        self.assertIn("厂区", judged["结论"])
        self.assertNotEqual(judged["分型倾向"], "D")

    def test_small_charge_is_red(self) -> None:
        row = _row(
            场站名称="小站",
            省="辽宁",
            市="沈阳",
            场景类型="纯充电",
            有充电="是",
            充电额定kW=60,
            变压器额定kVA=100,
            独立户号="是",
            能调控="能",
            场地能否放储能柜="是",
            用电性质="充电设施",
        )
        judged = screen_station(row, self.ln)
        self.assertEqual(judged["红绿灯"], "红")

    def test_existing_storage_is_gray(self) -> None:
        row = _row(
            场站名称="已有柜",
            有储能="是",
            储能kWh=262,
            储能功率kW=125,
        )
        judged = screen_station(row, self.ln)
        self.assertEqual(judged["红绿灯"], "灰")

    def test_liaoning_charge_weak_spread_not_green_by_piles_only(self) -> None:
        """辽宁充电站：有桩不够，价差贴近闸门 + 42号免需量 → 不配储。"""
        row = _row(
            场站名称="浑南充电站",
            省="辽宁",
            市="沈阳",
            场景类型="纯充电",
            有充电="是",
            充电额定kW=1200,
            充电可调="是",
            变压器额定kVA=4000,
            独立户号="是",
            能调控="能",
            场地能否放储能柜="是",
            用电性质="充电设施",
            日均用电或充电量kWh=4000,
            购电方式="电网代理购电",
        )
        judged = screen_station(row, self.ln)
        self.assertEqual(judged["红绿灯"], "红")
        self.assertEqual(judged["分型倾向"], "D")

    def test_guangdong_charge_can_be_green(self) -> None:
        row = _row(
            场站名称="广州超充",
            省="广东",
            市="广州",
            场景类型="纯充电",
            有充电="是",
            充电额定kW=1200,
            充电可调="是",
            变压器额定kVA=4000,
            独立户号="是",
            能调控="能",
            场地能否放储能柜="是",
            用电性质="充电设施",
            日均用电或充电量kWh=4000,
            购电方式="电网代理购电",
        )
        judged = screen_station(row, self.gd)
        self.assertEqual(judged["红绿灯"], "绿")
        self.assertIn("充电", judged["结论"])

    def test_storage_green_without_ctrl(self) -> None:
        """储能投资不看能否调控：桩不可调仍可绿。"""
        row = _row(
            场站名称="广州不可调桩",
            省="广东",
            市="广州",
            场景类型="纯充电",
            充电kW=1200,
            变压器额定kVA=4000,
            独立户号="是",
            能调控="不能",
            场地能否放储能柜="是",
            用电性质="充电设施",
            日均用电或充电量kWh=4000,
            购电方式="电网代理购电",
        )
        judged = screen_station(row, self.gd)
        self.assertEqual(judged["红绿灯"], "绿")

    def test_reason_explains_value_in_full_sentences(self) -> None:
        """原因和下一步必须是完整句子，讲清值或不值。"""
        green = screen_station(
            _row(
                场站名称="广州超充",
                省="广东",
                市="广州",
                场景类型="纯充电",
                有充电="是",
                充电额定kW=1200,
                充电可调="是",
                变压器额定kVA=4000,
                独立户号="是",
                能调控="能",
                场地能否放储能柜="是",
                用电性质="充电设施",
                日均用电或充电量kWh=4000,
                购电方式="电网代理购电",
            ),
            self.gd,
        )
        self.assertGreaterEqual(green["原因"].count("。"), 2)
        self.assertGreater(len(green["推荐下一步"]), 18)
        self.assertIn("价值", green["原因"])

        red = screen_station(
            _row(
                场站名称="浑南充电站",
                省="辽宁",
                市="沈阳",
                场景类型="纯充电",
                有充电="是",
                充电额定kW=1200,
                变压器额定kVA=4000,
                独立户号="是",
                能调控="能",
                场地能否放储能柜="是",
                用电性质="充电设施",
                日均用电或充电量kWh=4000,
                购电方式="电网代理购电",
            ),
            self.ln,
        )
        self.assertGreaterEqual(red["原因"].count("。"), 2)
        self.assertIn("价值", red["原因"])
        self.assertIn("不要为配储", red["推荐下一步"])

    def test_run_screen_batch(self) -> None:
        df = pd.DataFrame(
            [
                {
                    "场站名称": "铁西机械厂",
                    "省": "辽宁",
                    "市": "沈阳",
                    "场景类型": "厂区负荷",
                    "厂区白天最小kW": 380,
                    "其他可调kW": 80,
                    "变压器额定kVA": 800,
                    "独立户号": "是",
                    "共表或分表": "分表独立户号",
                    "能调控": "能",
                    "场地能否放储能柜": "是",
                    "用电性质": "工业",
                    "日均用电或充电量kWh": 9000,
                }
            ]
        )
        out = run_screen(df)
        self.assertEqual(len(out), 1)
        self.assertEqual(out.iloc[0]["电价库"], "辽宁（含沈阳）")
        self.assertGreater(float(out.iloc[0]["粗可调kW"]), 0)

    def test_vpp_legacy_columns_still_work(self) -> None:
        """三省 VPP 原表头：城市 / 充电kW / 独立户号? / 其他可调kW。"""
        df = pd.DataFrame(
            [
                {
                    "场站名称": "南山光储充站",
                    "城市": "深圳",
                    "区县+地址": "南山区xx",
                    "资源类型": "光储充",
                    "充电kW": 1200,
                    "储能功率kW": 0,
                    "储能kWh": 0,
                    "光伏kWp": 100,
                    "其他可调kW": 0,
                    "可持续h": 2,
                    "电费户号": "03xxxxxxxxxx",
                    "独立户号?": "独立",
                    "能调控?": "能",
                    "用电性质": "非工业",
                    "加密网关?": "未装",
                    "变压器额定kVA": 4000,
                    "日均用电或充电量kWh": 4000,
                    "场地能否放储能柜": "是",
                    "购电方式": "电网代理购电",
                }
            ]
        )
        out = run_screen(df)
        self.assertEqual(out.iloc[0]["市"], "深圳")
        self.assertEqual(out.iloc[0]["电价库"], "深圳供电局")
        self.assertAlmostEqual(float(out.iloc[0]["粗可调kW"]), 600.0)
        self.assertEqual(out.iloc[0]["红绿灯"], "绿")

    def test_renamed_headers_and_old_aliases(self) -> None:
        """新列名可筛；旧列名经 alias 仍能导入。"""
        new_df = pd.DataFrame(
            [
                {
                    "场站名称": "铁西新列名",
                    "省": "辽宁",
                    "市": "沈阳",
                    "资源类型": "厂区负荷",
                    "白天压不下去的负荷（千瓦）": 380,
                    "厂区其他可调（千瓦）": 80,
                    "变压器额定kVA": 800,
                    "独立户号?": "独立",
                    "能调控?": "能",
                    "场地能否放储能柜": "是",
                    "用电性质": "工业",
                    "日均用电或充电量kWh": 9000,
                    "购电方式": "电网代理购电",
                }
            ]
        )
        out_new = run_screen(new_df)
        self.assertIn(out_new.iloc[0]["红绿灯"], ("绿", "黄"))
        self.assertGreater(float(out_new.iloc[0]["粗可调kW"]), 0)

        old_df = pd.DataFrame(
            [
                {
                    "场站名称": "铁西旧列名",
                    "省": "辽宁",
                    "市": "沈阳",
                    "资源类型": "厂区负荷",
                    "厂区白天最小kW": 380,
                    "其他可调kW": 80,
                    "可持续h": 1,
                    "充电可调kW": 0,
                    "剩余可用kVA": 120,
                    "合同需量kW": 600,
                    "变压器额定kVA": 800,
                    "独立户号?": "独立",
                    "能调控?": "能",
                    "场地能否放储能柜": "是",
                    "用电性质": "工业",
                    "日均用电或充电量kWh": 9000,
                    "购电方式": "电网代理购电",
                }
            ]
        )
        out_old = run_screen(old_df)
        self.assertIn(out_old.iloc[0]["红绿灯"], ("绿", "黄"))
        self.assertGreater(float(out_old.iloc[0]["粗可调kW"]), 0)


class ScoreTests(unittest.TestCase):
    def test_factory_peak_shave_type_a(self) -> None:
        r = score_and_invest(
            ScoreInput(
                scene="厂区",
                tf_kva=800,
                peak_kw=850,
                overload_hours=1.0,
                valley_idle_hours=0.5,
                valley_headroom_kw=50,
                peak_yuan=0.82,
                valley_yuan=0.38,
                two_part=True,
                demand_waiver=False,
                demand_yuan_per_kw_month=32.0,
                daily_kwh=9000,
                service_fee=0.0,
                purchase_market=False,
                site_ok=True,
                chg_controllable=False,
            )
        )
        self.assertEqual(r.type_code, "A")
        self.assertGreaterEqual(r.n_cabinets, 1)

    def test_charge_waiver_not_type_a(self) -> None:
        r = score_and_invest(
            ScoreInput(
                scene="充电",
                tf_kva=4000,
                peak_kw=3600,
                overload_hours=1.0,
                valley_idle_hours=3.0,
                valley_headroom_kw=800,
                peak_yuan=1.10,
                valley_yuan=0.27,
                two_part=True,
                demand_waiver=True,
                demand_yuan_per_kw_month=36.0,
                daily_kwh=4000,
                service_fee=0.4,
                purchase_market=False,
                site_ok=True,
                chg_controllable=True,
            )
        )
        self.assertNotEqual(r.type_code, "A")
        self.assertIn(r.type_code, ("B", "C", "D", "E"))

    def test_hard_gate_all_fit_scores_100(self) -> None:
        from storage_scoring import HARD_ITEMS

        answers = {it["id"]: "适合" for it in HARD_ITEMS}
        gate = evaluate_hard_gate(answers)
        self.assertTrue(gate.complete)
        self.assertTrue(gate.passed)
        self.assertEqual(gate.score, 100)

    def test_hard_gate_one_unfit_blocks(self) -> None:
        from storage_scoring import HARD_ITEMS

        answers = {it["id"]: "适合" for it in HARD_ITEMS}
        answers["fire"] = "不适合"
        gate = evaluate_hard_gate(answers)
        self.assertTrue(gate.complete)
        self.assertFalse(gate.passed)
        self.assertIn("消防与防火间距", gate.fails)
        r = score_and_invest(
            ScoreInput(
                scene="厂区",
                tf_kva=800,
                peak_kw=850,
                overload_hours=1.0,
                valley_idle_hours=2.0,
                valley_headroom_kw=200,
                peak_yuan=0.82,
                valley_yuan=0.38,
                two_part=True,
                demand_waiver=False,
                demand_yuan_per_kw_month=32.0,
                daily_kwh=9000,
                service_fee=0.0,
                purchase_market=False,
                site_ok=False,
                chg_controllable=False,
                hard_passed=False,
                site_score_pts=gate.score,
            )
        )
        self.assertEqual(r.type_code, "D")
        self.assertEqual(r.n_cabinets, 0)

    def test_score_table_xlsx_matches_code(self) -> None:
        from openpyxl import load_workbook
        from storage_scoring import HARD_ITEMS

        path = Path(__file__).resolve().parent / "doc" / "储能配置打分表.xlsx"
        if not path.exists():
            self.skipTest("打分表不在 doc/")
        ws = load_workbook(path, data_only=True).active
        names = [row[1] for row in ws.iter_rows(min_row=2, max_col=3, values_only=True) if row[1]]
        self.assertEqual(names, [it["name"] for it in HARD_ITEMS])
        weights = [int(row[2]) for row in ws.iter_rows(min_row=2, max_col=3, values_only=True) if row[2] is not None]
        self.assertEqual(sum(weights), 100)


class ExcelReadTests(unittest.TestCase):
    def test_read_screening_sample(self) -> None:
        from ui_pipeline import _read_station_sheet

        path = Path(__file__).resolve().parent / "doc" / "测试实例_配储筛站摸底表.xlsx"
        if not path.exists():
            self.skipTest("测试实例尚未生成")
        class _Up:
            name = path.name

            def getvalue(self):
                return path.read_bytes()

        raw = _read_station_sheet(_Up())
        self.assertGreaterEqual(len(raw), 12)
        out = run_screen(raw)
        self.assertEqual(set(out["红绿灯"]), {"绿", "黄", "红", "灰"})
        self.assertTrue((out["原因"].astype(str).str.len() > 40).all())

        from io import BytesIO

        from openpyxl import load_workbook

        from intake_screening import order_screen_columns
        from ui_pipeline import _export_screen_xlsx

        ordered = order_screen_columns(out)
        self.assertEqual(list(ordered.columns)[:4], ["红绿灯", "结论", "推荐下一步", "原因"])
        data = _export_screen_xlsx(out)
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        self.assertEqual([c.value for c in ws[1][:4]], ["红绿灯", "结论", "推荐下一步", "原因"])
        yellow_a = str(ws.cell(2, 2).fill.fgColor.rgb or ws.cell(2, 2).fill.start_color.rgb)
        yellow_b = str(ws.cell(2, 3).fill.fgColor.rgb or ws.cell(2, 3).fill.start_color.rgb)
        self.assertIn("FFF3B0", yellow_a)
        self.assertIn("FFF3B0", yellow_b)


if __name__ == "__main__":
    unittest.main()
