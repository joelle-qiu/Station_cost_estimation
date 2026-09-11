# -*- coding: utf-8 -*-
"""生成对客统一摸底表：蓝列 VPP 组盘 + 绿列储能投资筛站（运营商/厂区共用）。作者: jiali.qiu"""

from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from intake_screening import (
    ACCT_OPTS,
    ARCH_OPTS,
    CTRL_OPTS,
    GATEWAY_OPTS,
    GENSET_OPTS,
    HOURS_OPTS,
    INTAKE_COLUMNS,
    PURCHASE,
    SCENE_OPTIONS,
    STORAGE_FIRST_COL,
    STORAGE_LAST_COL,
    USER_CLASS,
    VPP_LAST_COL,
    YES_NO,
)

ROOT = Path(__file__).resolve().parent
OUT_MAIN = ROOT / "doc" / "场站资源摸底收资表.xlsx"
OUT_ALIAS = ROOT / "doc" / "通用场站配储摸底收资表.xlsx"
VG = Path(r"c:\Users\joelle.qiu\Documents\Envision\CusorPath\VirtualGrid")
VG_COPIES = [
    VG / "product_vpp_operator_bp" / "场站资源摸底收资表.xlsx",
    VG / "outputs" / "场站资源摸底收资表.xlsx",
]

NAVY = "0A3D6B"
NAVY2 = "1E4A73"
STOR = "1A6B4A"
YELLOW = "FFF4C4"
VPP_INPUT = "FFF4C4"
STOR_INPUT = "C8EBD8"
HINT = "F4F6F9"
VPP_HINT = "E8EEF5"
STOR_HINT = "E3F4EA"
EX = "F1F5F9"
VPP_EX = "E2E8F0"
STOR_EX = "D5EADF"
NOTE_HEAD = "64748B"
NOTE_INPUT = "F1F5F9"
INK = "1E293B"
SLATE = "475569"
WHITE = "FFFFFF"
THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
WRAP = Alignment(wrap_text=True, vertical="center")
LEFT = Alignment(wrap_text=True, vertical="center", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
FONT_H = Font(name="微软雅黑", color=WHITE, bold=True, size=10)
FONT_N = Font(name="微软雅黑", size=10, color=INK)

FIRST, LAST = 8, 200  # 第 4 行填法，5–7 示例，从第 8 行正式填

COL_HINTS = {
    "场站名称": "请勿填写本行，下面从示例之后开始填",
    "省": "与电费单一致",
    "市": "深圳必须单独写市",
    "区县+地址": "能找到这个站即可",
    "资源类型": "下拉。无桩厂区选「厂区负荷」或「可调负荷」",
    "充电kW": "本站充电桩额定合计。没有桩填 0",
    "储能功率kW": "已装储能的功率。没有填 0",
    "储能kWh": "已装储能的容量。没有填 0",
    "光伏kWp": "没有填 0",
    "厂区其他可调（千瓦）": "空调、产线等，不要把充电桩算进来。没有填 0",
    "一次可调控时长（小时）": "一次降负荷大概能维持多久，不是电池能用几年",
    "电费户号": "电费发票上的营销户号，不是充电 App 桩号",
    "独立户号?": "是否单独和电网结算。和物业合用一块表选「合表」",
    "能调控?": "这一站能不能远程调节。公司整体情况请填第 1 页",
    "用电性质": "充电站选「充电设施」，不要选「非工业」。工厂选「工业」，商场选「商业」",
    "加密网关?": "仅深圳站填已装或未装，其他城市选「不适用」",
    "充电桩型": "无桩选「不适用」",
    "充电桩可远程限功率（千瓦）": "能远程把功率降下来的桩有多少千瓦。不能控填 0",
    "V2G可调kW": "车辆能向电网放电的功率。没有填 0",
    "换电仓功率能否远程下调": "问功率能不能远程往下调。没有换电站选「不适用」",
    "柴发": "问能不能调度。没有柴油发电机选「无」",
    "光伏可限发kW": "能远程限制发出功率的千瓦。不能调填 0",
    "白天压不下去的负荷（千瓦）": "白天一直在用、减不掉的电。纯充电站填 0。不是「能调多少」",
    "变压器额定kVA": "变压器铭牌容量",
    "变压器还能再接（千伏安）": "变压器还能再接多少设备。不清楚可空着",
    "电费合同需量（千瓦）": "电费合同里约定的需量。没有这项填 0",
    "日均用电或充电量kWh": "近月日均。充电站填充电量，厂区填用电量",
    "场地能否放储能柜": "柜子大约一个集装箱大小。放不下选「否」",
    "购电方式": "还在电网买电选「电网代理购电」；已找售电公司或直接入市选另一项",
    "备注": "选填",
}


def _font(size=10, bold=False, color=INK):
    return Font(name="微软雅黑", size=size, bold=bold, color=color)


def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def _dv(ws, opts: list[str], cells: str, title: str, prompt: str) -> None:
    formula = '"' + ",".join(opts) + '"'
    dv = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="请从下拉选",
        error="请从下拉列表选择。",
        promptTitle=title,
        prompt=prompt,
        showInputMessage=True,
    )
    ws.add_data_validation(dv)
    dv.add(cells)


def _blank_row() -> list:
    return [""] * len(INTAKE_COLUMNS)


def _ex(*pairs) -> list:
    row = _blank_row()
    idx = {n: i for i, n in enumerate(INTAKE_COLUMNS)}
    for k, v in pairs:
        row[idx[k]] = v
    return row


def build() -> Path:
    wb = Workbook()
    col = {name: i + 1 for i, name in enumerate(INTAKE_COLUMNS)}
    last_letter = get_column_letter(len(INTAKE_COLUMNS))

    ws0 = wb.active
    ws0.title = "0-请先看"
    ws0["A1"] = "场站资源摸底收资表"
    ws0["A1"].font = Font(name="微软雅黑", bold=True, size=16, color=WHITE)
    ws0["A1"].fill = _fill(NAVY)
    ws0.merge_cells("A1:B1")
    ws0.row_dimensions[1].height = 28
    blocks = [
        (
            "这张表做什么",
            "请把各场站的基本情况填在同一张表里。我们用来做两件事：看哪些站可以参加虚拟电厂，以及看哪些站值得进一步评估是否配置储能。"
            "您不必填写电价和投资回报，现在也不必上传用电曲线。",
        ),
        (
            "请填哪几页",
            "请填写「1-公司概况」和「2-资源清单」。其余工作表由我们内部使用，不用填、也不用发给其他人。",
        ),
        (
            "「2-资源清单」怎么填",
            "一行只写一个场站，不要把全市设备加总写在一行。没有的设备在对应格子填 0。"
            "充电站、工厂、园区都用这一页；没有充电桩也可以填，把充电相关格子填 0 即可。"
            "表头下一行是填法说明，可对照填写。",
        ),
        (
            "蓝色列填写什么",
            "请按场站填写充电桩、储能、光伏、厂区其他可调负荷的容量，以及电费户号、是否单独立户、这一站能不能远程调控。"
            "换电仓、柴油发电机、充电桩型也在蓝列：问的是能不能远程调度，不是有没有这些设备。",
        ),
        (
            "绿色列填写什么",
            "请填写变压器容量、每天用电或充电量、场地能不能放置储能柜、以及现在如何购电。"
            "纯充电站把「白天压不下去的负荷」填 0。绿色列不问设备能不能调控。",
        ),
        (
            "您不用填写的内容",
            "不用填电网节点编号、峰谷电价、投资回收期。地址和户号填写后，由我们内部核对。"
            "如果后续需要用电曲线，我们会再向需要评估储能的场站单独索取。",
        ),
    ]
    r = 3
    block_fill = {
        "蓝色列填写什么": NAVY2,
        "绿色列填写什么": STOR,
    }
    for h, b in blocks:
        ws0.merge_cells(f"A{r}:B{r}")
        ws0[f"A{r}"] = h
        ws0[f"A{r}"].font = _font(11, True, WHITE)
        ws0[f"A{r}"].fill = _fill(block_fill.get(h, NAVY2))
        r += 1
        ws0.merge_cells(f"A{r}:B{r}")
        ws0[f"A{r}"] = b
        ws0[f"A{r}"].font = FONT_N
        body_fill = VPP_HINT if h.startswith("蓝色") else (STOR_HINT if h.startswith("绿色") else HINT)
        ws0[f"A{r}"].fill = _fill(body_fill)
        ws0[f"A{r}"].alignment = WRAP
        ws0.row_dimensions[r].height = 64
        r += 1
    ws0.column_dimensions["A"].width = 22
    ws0.column_dimensions["B"].width = 96

    ws1 = wb.create_sheet("1-公司概况")
    ws1.merge_cells("A1:D1")
    ws1["A1"] = "1 公司概况（约2分钟）"
    ws1["A1"].font = _font(16, True, WHITE)
    ws1["A1"].fill = _fill(NAVY)
    ws1.merge_cells("A2:D2")
    ws1["A2"] = "黄色格子填写或下拉。充电网络、产业园、工厂都可以填。没有充电桩也请写明（清单里对应格子填 0）。资源规模请按充电、储能、光伏、厂区分开写，不要只写一个总数。"
    ws1["A2"].fill = _fill(HINT)
    for i, t in enumerate(["序号", "问题", "请填写", "说明"], 1):
        c = ws1.cell(3, i, t)
        c.fill = _fill(NAVY2)
        c.font = FONT_H
        c.alignment = CENTER
        c.border = THIN
    questions = [
        (1, "公司名称", "营业执照名称即可"),
        (2, "联系人 / 电话", "商务 + 运维一人即可"),
        (3, "资产类型", "充电网络 / 厂区或园区 / 两者都有"),
        (4, "各类资源大约多少", "例：充电 8MW / 储能 2MW·4MWh / 光伏 0.5MWp / 厂区可调 1MW"),
        (5, "主要在哪几个省、市", "例：辽宁沈阳、广东深圳。深圳必须单独标"),
        (6, "现在电费跟哪家售电公司签", "没有就填「没有」"),
        (7, "有没有签过虚拟电厂 / 需求响应", "含国网综能、南网电动、深电能、绿色国网"),
        (8, "公司整体是否具备远程调控能力", "这里只填公司整体情况。每一站能控多少千瓦，请到「2-资源清单」分站填写。"),
        (9, "深圳站能否配合安装加密网关", "没有深圳站选「不适用」"),
        (10, "是否已在「绿色国网」申报过", "仅辽宁需求响应用；不是辽宁选「不适用」"),
        (11, "自有平台能否导出功率曲线", "现在不用上传。如需曲线我们会再联系。"),
    ]
    for i, (n, q, help_txt) in enumerate(questions):
        rr = 4 + i
        ws1.cell(rr, 1, n).fill = _fill(NAVY2)
        ws1.cell(rr, 1).font = _font(10, True, WHITE)
        ws1.cell(rr, 1).alignment = CENTER
        ws1.cell(rr, 2, q).fill = _fill(HINT)
        ws1.cell(rr, 3).fill = _fill(YELLOW)
        ws1.cell(rr, 4, help_txt).fill = _fill(HINT)
        ws1.cell(rr, 4).font = _font(9, False, SLATE)
        for cc in range(1, 5):
            ws1.cell(rr, cc).border = THIN
            ws1.cell(rr, cc).alignment = WRAP
        ws1.row_dimensions[rr].height = 28
    for coln, w in zip("ABCD", (8, 38, 42, 48)):
        ws1.column_dimensions[coln].width = w
    _dv(ws1, ["充电网络", "厂区或园区", "两者都有"], "C6", "资产", "可先选一类")
    _dv(ws1, ["没有", "有（未到期）", "有（快到期）", "不确定"], "C10", "绑定", "合同期内原则上一家 VPP")
    _dv(ws1, ["都能", "仅储能", "仅充电", "部分能", "都不能", "不确定"], "C11", "调控", "公司整体情况；分站千瓦请填第 2 页")
    _dv(ws1, ["可以", "需评估", "不可以", "不适用（无深圳站）", "不确定"], "C12", "网关", "深圳精准响应四步")
    _dv(ws1, ["已申报", "没有", "不适用", "不确定"], "C13", "绿色国网", "辽宁 181 号")
    _dv(ws1, ["能", "否", "需对接", "不确定"], "C14", "曲线", "现在不用上传")

    ws2 = wb.create_sheet("2-资源清单")
    vpp_last = col[VPP_LAST_COL]
    stor_first = col[STORAGE_FIRST_COL]
    stor_last = col[STORAGE_LAST_COL]
    note_col = col["备注"]

    def _zone(cidx: int) -> str:
        if cidx <= vpp_last:
            return "vpp"
        if cidx <= stor_last:
            return "stor"
        return "note"

    def _zone_head(cidx: int) -> str:
        return {"vpp": NAVY2, "stor": STOR, "note": NOTE_HEAD}[_zone(cidx)]

    def _zone_input(cidx: int) -> str:
        return {"vpp": VPP_INPUT, "stor": STOR_INPUT, "note": NOTE_INPUT}[_zone(cidx)]

    def _zone_ex(cidx: int) -> str:
        return {"vpp": VPP_EX, "stor": STOR_EX, "note": EX}[_zone(cidx)]

    ws2.merge_cells(f"A1:{last_letter}1")
    ws2["A1"] = "2 资源清单（一行只写一个场站）  蓝色列：装了什么、户号、能不能调度    绿色列：变压器、电量、场地、如何购电"
    ws2["A1"].font = _font(16, True, WHITE)
    ws2["A1"].fill = _fill(NAVY)
    ws2.merge_cells(start_row=2, start_column=1, end_row=2, end_column=vpp_last)
    ws2.cell(2, 1, "① 蓝色列：请填写本站装了什么、电费户号、能不能远程调度。没有的设备填 0。")
    ws2.cell(2, 1).font = _font(10, True, WHITE)
    ws2.cell(2, 1).fill = _fill(NAVY2)
    ws2.cell(2, 1).alignment = WRAP
    for cidx in range(1, vpp_last + 1):
        ws2.cell(2, cidx).fill = _fill(NAVY2)
        ws2.cell(2, cidx).border = THIN
    ws2.merge_cells(start_row=2, start_column=stor_first, end_row=2, end_column=stor_last)
    ws2.cell(
        2,
        stor_first,
        "② 绿色列：请填写变压器、每天用电或充电量、场地能不能放储能柜、如何购电。不问能不能调控。纯充电站把「白天压不下去的负荷」填 0。",
    )
    ws2.cell(2, stor_first).font = _font(10, True, WHITE)
    ws2.cell(2, stor_first).fill = _fill(STOR)
    ws2.cell(2, stor_first).alignment = WRAP
    for cidx in range(stor_first, stor_last + 1):
        ws2.cell(2, cidx).fill = _fill(STOR)
        ws2.cell(2, cidx).border = THIN
    ws2.cell(2, note_col, "共用")
    ws2.cell(2, note_col).font = _font(10, True, WHITE)
    ws2.cell(2, note_col).fill = _fill(NOTE_HEAD)
    ws2.cell(2, note_col).alignment = CENTER
    ws2.cell(2, note_col).border = THIN
    ws2.row_dimensions[2].height = 36
    for i, name in enumerate(INTAKE_COLUMNS, 1):
        cell = ws2.cell(3, i, name)
        cell.fill = _fill(_zone_head(i))
        cell.font = FONT_H
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    ws2.row_dimensions[3].height = 40

    for i, name in enumerate(INTAKE_COLUMNS, 1):
        cell = ws2.cell(4, i, COL_HINTS.get(name, ""))
        cell.fill = _fill({"vpp": VPP_HINT, "stor": STOR_HINT, "note": HINT}[_zone(i)])
        cell.font = _font(8, False, SLATE)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        cell.border = THIN
    ws2.row_dimensions[4].height = 56

    examples = [
        _ex(
            ("场站名称", "示例：浑南充电站"),
            ("省", "辽宁"),
            ("市", "沈阳"),
            ("区县+地址", "浑南区xx路"),
            ("资源类型", "纯充电"),
            ("充电kW", 1200),
            ("储能功率kW", 0),
            ("储能kWh", 0),
            ("光伏kWp", 0),
            ("厂区其他可调（千瓦）", 0),
            ("一次可调控时长（小时）", "2"),
            ("电费户号", "21xxxxxxxxxx"),
            ("独立户号?", "独立"),
            ("能调控?", "能"),
            ("用电性质", "充电设施"),
            ("加密网关?", "不适用"),
            ("光伏可限发kW", 0),
            ("充电桩型", "分体式"),
            ("充电桩可远程限功率（千瓦）", 600),
            ("V2G可调kW", 0),
            ("换电仓功率能否远程下调", "不适用"),
            ("柴发", "无"),
            ("白天压不下去的负荷（千瓦）", 0),
            ("变压器额定kVA", 4000),
            ("变压器还能再接（千伏安）", 200),
            ("电费合同需量（千瓦）", 0),
            ("日均用电或充电量kWh", 4000),
            ("场地能否放储能柜", "是"),
            ("购电方式", "电网代理购电"),
            ("备注", "这是示例，正式填写请删除本行。充电站：白天负荷填 0，用电性质选充电设施。"),
        ),
        _ex(
            ("场站名称", "示例：宝安物流园配电"),
            ("省", "广东"),
            ("市", "深圳"),
            ("区县+地址", "宝安区xx园"),
            ("资源类型", "厂区+充电"),
            ("充电kW", 480),
            ("储能功率kW", 0),
            ("储能kWh", 0),
            ("光伏kWp", 0),
            ("厂区其他可调（千瓦）", 120),
            ("一次可调控时长（小时）", "1"),
            ("电费户号", "03xxxxxxxxxx"),
            ("独立户号?", "独立"),
            ("能调控?", "需改造"),
            ("用电性质", "工业"),
            ("加密网关?", "未装"),
            ("光伏可限发kW", 0),
            ("充电桩型", "一体式"),
            ("充电桩可远程限功率（千瓦）", 0),
            ("V2G可调kW", 0),
            ("换电仓功率能否远程下调", "不适用"),
            ("柴发", "无"),
            ("白天压不下去的负荷（千瓦）", 350),
            ("变压器额定kVA", 1000),
            ("变压器还能再接（千伏安）", 150),
            ("电费合同需量（千瓦）", 800),
            ("日均用电或充电量kWh", 12000),
            ("场地能否放储能柜", "是"),
            ("购电方式", "电网代理购电"),
            ("备注", "这是示例，正式填写请删除本行。厂区带充电：绿色列填变压器和白天负荷，不问桩能不能调。"),
        ),
        _ex(
            ("场站名称", "示例：铁西机械厂配电"),
            ("省", "辽宁"),
            ("市", "沈阳"),
            ("区县+地址", "铁西区xx路"),
            ("资源类型", "可调负荷"),
            ("充电kW", 0),
            ("储能功率kW", 0),
            ("储能kWh", 0),
            ("光伏kWp", 0),
            ("厂区其他可调（千瓦）", 80),
            ("一次可调控时长（小时）", "1"),
            ("电费户号", "21xxxxxxxxxx"),
            ("独立户号?", "独立"),
            ("能调控?", "能"),
            ("用电性质", "工业"),
            ("加密网关?", "不适用"),
            ("光伏可限发kW", 0),
            ("充电桩型", "不适用"),
            ("充电桩可远程限功率（千瓦）", 0),
            ("V2G可调kW", 0),
            ("换电仓功率能否远程下调", "不适用"),
            ("柴发", "无"),
            ("白天压不下去的负荷（千瓦）", 380),
            ("变压器额定kVA", 800),
            ("变压器还能再接（千伏安）", 120),
            ("电费合同需量（千瓦）", 600),
            ("日均用电或充电量kWh", 9000),
            ("场地能否放储能柜", "是"),
            ("购电方式", "电网代理购电"),
            ("备注", "这是示例，正式填写请删除本行。无桩厂区：充电填 0、桩型选不适用，绿色列仍填变压器和日电量。"),
        ),
    ]
    for i, ex in enumerate(examples):
        rr = 5 + i
        for c, val in enumerate(ex, 1):
            cell = ws2.cell(rr, c, val)
            cell.fill = _fill(_zone_ex(c))
            cell.font = _font(9, False, SLATE)
            cell.border = THIN
            cell.alignment = WRAP
    for rr in range(FIRST, min(FIRST + 14, LAST) + 1):
        for c in range(1, len(INTAKE_COLUMNS) + 1):
            cell = ws2.cell(rr, c)
            cell.fill = _fill(_zone_input(c))
            cell.border = THIN
    for i, _ in enumerate(INTAKE_COLUMNS, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = 12
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["D"].width = 20
    for name, w in (
        ("厂区其他可调（千瓦）", 16),
        ("一次可调控时长（小时）", 16),
        ("充电桩可远程限功率（千瓦）", 18),
        ("换电仓功率能否远程下调", 16),
        ("白天压不下去的负荷（千瓦）", 18),
        ("变压器还能再接（千伏安）", 16),
        ("电费合同需量（千瓦）", 16),
        ("日均用电或充电量kWh", 14),
        ("用电性质", 12),
    ):
        ws2.column_dimensions[get_column_letter(col[name])].width = w
    ws2.freeze_panes = "B5"
    ws2.auto_filter.ref = f"A3:{last_letter}{LAST}"
    split_border = Border(
        left=Side(style="medium", color=STOR),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    for rr in range(2, min(FIRST + 14, LAST) + 1):
        ws2.cell(rr, stor_first).border = split_border

    def rng(key: str) -> str:
        letter = get_column_letter(col[key])
        return f"{letter}{FIRST}:{letter}{LAST}"

    _dv(ws2, SCENE_OPTIONS, rng("资源类型"), "类型", "无桩厂区选「厂区负荷」或「可调负荷」，绿色列仍填变压器和电量")
    _dv(ws2, HOURS_OPTS, rng("一次可调控时长（小时）"), "时长", "一次降负荷大概能维持多久，不是电池寿命")
    _dv(ws2, ACCT_OPTS, rng("独立户号?"), "户号", "和物业共用一块电表请选合表")
    _dv(ws2, CTRL_OPTS, rng("能调控?"), "调控", "这一站能不能远程调节。绿色列不看这一列")
    _dv(ws2, USER_CLASS, rng("用电性质"), "性质", "充电站选充电设施，不要选非工业。工厂选工业，商场选商业")
    _dv(ws2, GATEWAY_OPTS, rng("加密网关?"), "网关", "仅深圳填已装或未装，其他城市选不适用")
    _dv(ws2, ARCH_OPTS, rng("充电桩型"), "桩型", "无桩选不适用")
    _dv(ws2, YES_NO + ["不适用"], rng("换电仓功率能否远程下调"), "换电仓", "问功率能不能远程往下调。没有换电站选不适用")
    _dv(ws2, GENSET_OPTS, rng("柴发"), "柴发", "问能不能调度。没有选无")
    _dv(ws2, YES_NO, rng("场地能否放储能柜"), "场地", "柜子大约集装箱大小。放不下选否")
    _dv(ws2, PURCHASE, rng("购电方式"), "购电", "还在电网买电选电网代理购电")

    # 3-内部评估：VPP 自动汇总（客户不填）
    st, co = "'2-资源清单'", "'1-公司概况'"
    ws3 = wb.create_sheet("3-内部评估")
    ws3.merge_cells("A1:D1")
    ws3["A1"] = "3 内部评估（客户不用填 · VPP 自动汇总；配储红绿灯在测算工具①）"
    ws3["A1"].font = _font(14, True, WHITE)
    ws3["A1"].fill = _fill(NAVY)
    ws3.merge_cells("A2:D2")
    ws3["A2"] = (
        "粗可调 = 充电×50% + 储能功率×90% + 其他可调×70%。光伏不计入负荷调节。"
        "辽宁 DR 聚合商 ≥5MW；广东全省 DR 同地市 ≥5MW；深圳精准响应无明文 MW。"
        "配储高价值站请把本文件导入测算工具页签①。"
    )
    ws3["A2"].fill = _fill(HINT)
    ws3["A2"].alignment = WRAP
    ws3.row_dimensions[2].height = 36
    for i, t in enumerate(["指标", "数值", "评分", "说明"], 1):
        c = ws3.cell(3, i, t)
        c.fill = _fill(NAVY2)
        c.font = FONT_H
        c.border = THIN

    def colL(name: str) -> str:
        return get_column_letter(col[name])

    metrics = [
        ("填报场站数（不含示例）", f"=COUNTA({st}!A{FIRST}:A{LAST})",
         "=IF(B4>=8,5,IF(B4>=5,4,IF(B4>=3,3,IF(B4>=1,2,1))))", "清单有效行"),
        ("其中沈阳站数", f'=COUNTIF({st}!C{FIRST}:C{LAST},"沈阳")',
         '=IF(B4=0,1,IF(B5/B4>=0.5,5,IF(B5>=3,4,IF(B5>=1,3,2))))', "辽宁示范密度"),
        ("其中深圳站数", f'=COUNTIF({st}!C{FIRST}:C{LAST},"深圳")',
         '=IF(B4=0,1,IF(B6/B4>=0.5,5,IF(B6>=3,4,IF(B6>=1,3,2))))', "深圳精准响应"),
        ("充电额定合计 kW", f"=SUM({st}!{colL('充电kW')}{FIRST}:{colL('充电kW')}{LAST})",
         "=IF(B7>=10000,5,IF(B7>=5000,4,IF(B7>=2000,3,IF(B7>=800,2,1))))", "负荷类"),
        ("储能功率合计 kW", f"=SUM({st}!{colL('储能功率kW')}{FIRST}:{colL('储能功率kW')}{LAST})",
         "=IF(B8>=2000,5,IF(B8>=500,4,IF(B8>=100,3,IF(B8>0,2,1))))", "储能类可单列"),
        ("光伏装机合计 kWp", f"=SUM({st}!{colL('光伏kWp')}{FIRST}:{colL('光伏kWp')}{LAST})",
         "=IF(B9>=1000,5,IF(B9>=200,4,IF(B9>0,3,3)))", "发电类另组；0 不扣分"),
        ("其他可调合计 kW", f"=SUM({st}!{colL('厂区其他可调（千瓦）')}{FIRST}:{colL('厂区其他可调（千瓦）')}{LAST})",
         "=IF(B10>=2000,5,IF(B10>=500,4,IF(B10>=100,3,IF(B10>0,2,1))))", "厂区/空调/产线"),
        ("粗可调 kW（负荷+储能）", "=B7*0.5+B8*0.9+B10*0.7",
         "=IF(B11>=5000,5,IF(B11>=2000,4,IF(B11>=1000,3,IF(B11>=500,2,1))))", "DR ≥5MW；现货单元实践 ≥1MW"),
        ("独立户号站数", f'=COUNTIF({st}!{colL("独立户号?")}{FIRST}:{colL("独立户号?")}{LAST},"独立")',
         "=IF(B4=0,1,IF(B12/B4>=0.8,5,IF(B12/B4>=0.5,3,2)))", "辽宁名录硬门槛"),
        ("能调控站数", f'=COUNTIF({st}!{colL("能调控?")}{FIRST}:{colL("能调控?")}{LAST},"能")+COUNTIF({st}!{colL("能调控?")}{FIRST}:{colL("能调控?")}{LAST},"部分能")+COUNTIF({st}!{colL("能调控?")}{FIRST}:{colL("能调控?")}{LAST},"需改造")',
         "=IF(B4=0,1,IF(B13/B4>=0.8,5,IF(B13/B4>=0.5,3,2)))", "含需改造；只服务 VPP"),
        ("已装加密网关站数", f'=COUNTIF({st}!{colL("加密网关?")}{FIRST}:{colL("加密网关?")}{LAST},"已装")',
         '=IF(B6=0,3,IF(B14>=B6,5,IF(B14>0,3,2)))', "仅深圳计"),
        ("售电/VPP 绑定", f"={co}!C10",
         '=IF(OR(B15="没有",B15=""),5,IF(B15="有（快到期）",4,IF(B15="不确定",3,1)))', "未到期=1"),
    ]
    for i, (name, formula, score, note) in enumerate(metrics):
        rr = 4 + i
        ws3.cell(rr, 1, name).border = THIN
        ws3.cell(rr, 2, formula).border = THIN
        ws3.cell(rr, 3, score).border = THIN
        ws3.cell(rr, 4, note).border = THIN
        ws3.cell(rr, 4).font = _font(9, False, SLATE)
    vr = 16
    ws3.cell(vr, 1, "综合均分").font = _font(10, True)
    ws3.cell(vr, 2, "=AVERAGE(C4:C15)")
    ws3.cell(vr, 3, '=IF(B16>=4,"建议 VPP/DR 试点",IF(B16>=3,"可培育",IF(B16>=2,"信息不足/暂缓","暂缓")))')
    ws3.cell(vr, 4, "配储绿/黄/红请导入测算工具①，不要用本页代替。")
    for cc in range(1, 5):
        ws3.cell(vr, cc).fill = _fill("FEF3C7")
        ws3.cell(vr, cc).border = THIN
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 22
    ws3.column_dimensions["C"].width = 28
    ws3.column_dimensions["D"].width = 40

    ws4 = wb.create_sheet("4-字段说明")
    tips = [
        ["区块", "字段", "谁用", "怎么填"],
        ["蓝", "资源类型", "双方", "纯充电、充电+储能、光储充、独立储能、分布式光伏、换电、可调负荷、厂区负荷、厂区+充电。无桩厂区选后两项。"],
        ["蓝", "充电kW 到厂区其他可调", "虚拟电厂", "按本站铭牌合计，没有填 0。不要把全市加总写在一行。无桩时充电填 0。厂区其他可调不要把充电桩算进去。"],
        ["蓝", "一次可调控时长（小时）", "虚拟电厂", "一次降负荷大概能维持多久，不是电池能用几年。常见不少于 1 小时。"],
        ["蓝", "独立户号?", "双方", "是否单独和电网结算。和物业合表时，虚拟电厂名录和储能需量账户都较难独立处理。"],
        ["蓝", "能调控? / 充电桩可远程限功率 / V2G", "仅虚拟电厂", "储能投资不看这些列。能调控指这一站能不能远程调节；具体千瓦填在充电桩可远程限功率或厂区其他可调。"],
        ["蓝", "用电性质", "双方", "充电站选「充电设施」，不要选「非工业」。工厂选「工业」，商场选「商业」。"],
        ["蓝", "加密网关?", "深圳虚拟电厂", "深圳站填已装或未装；其他城市选「不适用」。"],
        ["蓝", "充电桩型", "虚拟电厂 / 落地控制", "无桩选「不适用」。绿色列不问桩型。"],
        ["蓝", "换电仓功率能否远程下调", "仅虚拟电厂", "问的是仓位充电功率能不能远程往下调。没有换电站选「不适用」。"],
        ["蓝", "柴发", "仅虚拟电厂", "问能不能调度，不是有没有柴油机本身。绿色列不问。"],
        ["蓝", "光伏可限发kW", "仅虚拟电厂", "能远程限制发出功率的千瓦。不能调填 0。"],
        ["绿", "白天压不下去的负荷（千瓦）", "储能投资", "白天一直在用、减不掉的电，不是能调下来的负荷。产业园和工厂必填。纯充电站填 0。"],
        ["绿", "变压器 / 还能再接 / 合同需量 / 日均电量", "储能投资", "筛站看这些数。空着时该站最多进入需要补数的名单。峰负荷、超载时长、谷段空窗不在本表填写。"],
        ["绿", "场地 / 购电方式", "储能投资", "场地放不下储能柜则不配置储能。还在电网买电选「电网代理购电」；已找售电公司或直接入市选另一项。"],
        ["—", "不用填", "—", "电网节点编号、峰谷电价、投资回收期。城市用于匹配电价，户号用于内部核对。"],
    ]
    for row in tips:
        ws4.append(row)
    for cell in ws4[1]:
        cell.fill = _fill(NAVY2)
        cell.font = FONT_H
    for rr in range(2, ws4.max_row + 1):
        zone = str(ws4.cell(rr, 1).value or "")
        fill = VPP_HINT if zone.startswith("蓝") else (STOR_HINT if zone.startswith("绿") else HINT)
        head = NAVY2 if zone.startswith("蓝") else (STOR if zone.startswith("绿") else NOTE_HEAD)
        ws4.cell(rr, 1).fill = _fill(head)
        ws4.cell(rr, 1).font = _font(9, True, WHITE)
        for cc in range(2, 5):
            ws4.cell(rr, cc).fill = _fill(fill)
            ws4.cell(rr, cc).font = FONT_N
            ws4.cell(rr, cc).alignment = WRAP
        ws4.row_dimensions[rr].height = 36
    ws4.column_dimensions["A"].width = 16
    ws4.column_dimensions["B"].width = 36
    ws4.column_dimensions["C"].width = 22
    ws4.column_dimensions["D"].width = 88

    OUT_MAIN.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_MAIN)
    shutil.copy2(OUT_MAIN, OUT_ALIAS)
    for dest in VG_COPIES:
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(OUT_MAIN, dest)
    return OUT_MAIN


if __name__ == "__main__":
    print(build())
